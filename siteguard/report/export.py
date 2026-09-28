"""report：研判报告导出（JSON / Markdown / docx）。

M3：导出前由 kg/anchors.attach_anchors 为每条 clause_ref 注入 kg_node/anchor
字段——JSON 携带机器可读锚点，Markdown 的"依据条款"列渲染为 kg.html 可点击链接。

M5：导出前统一 sanitize_paths——仓库内绝对路径转 posix 相对路径（脱敏第③步：
本机路径不入产物）；新增 export_docx（python-docx optional，照 M2 docx 解析
容错范式：缺依赖抛 ValueError 带安装提示），报告含签署栏（编制/审核/批准 +
日期占位线）；docx 不含任何图片，天然 0 外链。
"""

import json
from pathlib import Path

from .. import config
from ..kg import anchors


def _clause_md(ref):
    """条款引用 → Markdown 片段：带锚点渲染为链接（M3），否则纯文本。"""
    label = "%s %s" % (ref.get("standard", ""), ref.get("clause", ""))
    anchor = ref.get("anchor")
    if anchor:
        return "[%s](%s)" % (label, anchor)
    return label


def sanitize_paths(report, root=None):
    """就地规范化报告中的本机路径：仓库内绝对路径 → posix 相对路径。

    覆盖 report["source_file"]、每张卡的 source["file"]、report["vision_images"]
    （脱敏第③步：个人/本机路径一律不入报告产物）。仓库外路径（测试临时目录、
    用户上传原始文件名）原样保留。返回 report 本身（便于链式使用）。
    """
    def conv(s):
        return config.display_path(s, root=root)

    if report.get("source_file"):
        report["source_file"] = conv(report["source_file"])
    for card in report.get("cards", []):
        src = card.get("source") or {}
        if src.get("file"):
            src["file"] = conv(src["file"])
    if report.get("vision_images"):
        report["vision_images"] = [conv(v) for v in report["vision_images"]]
    return report


def export_json(report, out_path):
    p = Path(out_path)
    p.parent.mkdir(parents=True, exist_ok=True)
    anchors.attach_anchors(report)
    sanitize_paths(report)
    with open(str(p), "w", encoding="utf-8") as f:
        json.dump(report, f, ensure_ascii=False, indent=2)
    return p


def export_md(report, out_path):
    p = Path(out_path)
    p.parent.mkdir(parents=True, exist_ok=True)
    anchors.attach_anchors(report)
    sanitize_paths(report)
    lines = []
    lines.append("# 安全风险隐患研判报告")
    lines.append("")
    lines.append("- 源文件：`%s`" % report.get("source_file", ""))
    lines.append("- 记录数：%d，识别隐患：%d" % (report.get("record_count", 0),
                                               report.get("hazard_count", 0)))
    lines.append("- 生成：SiteGuard v%s（LLM 兜底：%s）" % (
        report.get("version", "?"),
        "关闭" if report.get("settings", {}).get("disable_llm") else "开启"))
    lines.append("")
    lines.append("## 隐患清单")
    lines.append("")
    lines.append("| # | 原文摘录 | 隐患类型 | 类别 | 风险等级 | 依据条款 | 整改建议 |")
    lines.append("| --- | --- | --- | --- | --- | --- | --- |")
    for i, card in enumerate(report.get("cards", []), 1):
        j = card.get("judgment", {})
        clauses = "、".join(_clause_md(r) for r in j.get("clause_refs", [])) or "—"
        sugg = "；".join(s["text"] for s in j.get("suggestions", [])) or "—"
        lines.append("| %d | %s | %s | %s | **%s** | %s | %s |" % (
            i, card["quote"], card["hazard_type"], card["category"],
            j.get("risk_level", "—"), clauses, sugg))
    confirm = [c for c in report.get("cards", []) if c.get("judgment", {}).get("human_confirm")]
    fallback_cards = [c for c in report.get("cards", []) if c.get("fallback_clauses")]
    if fallback_cards:
        lines.append("")
        lines.append("## 兜底候选条款（向量检索，提示性信息，非判定依据）")
        lines.append("")
        for c in fallback_cards:
            cands = "、".join("%s %s（相似度 %.4f）" % (x["standard"], x["clause"], x["score"])
                              for x in c["fallback_clauses"])
            lines.append("- %s → %s" % (c["quote"], cands))
    if confirm:
        lines.append("")
        lines.append("## 待人工确认")
        lines.append("")
        for c in confirm:
            lines.append("- %s（%s）" % (c["quote"], c["judgment"]["hazard_type"]))
    if report.get("warnings"):
        lines.append("")
        lines.append("## 警告")
        lines.append("")
        for w in report["warnings"]:
            lines.append("- %s" % w)
    lines.append("")
    lines.append("> 本报告由规则引擎生成，风险等级与条款依据可溯源；种子条款以官方文本为准。")
    with open(str(p), "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return p


def export_docx(report, out_path):
    """研判报告 → Word docx（python-docx optional，缺依赖抛 ValueError 带安装提示）。

    版式：标题 → 元信息 → 隐患清单表 → 待人工确认 → 警告 → 结语 → 签署栏
    （编制/审核/批准 + 日期占位线）。docx 不含任何图片，天然 0 外链。
    """
    try:
        import docx  # python-docx（可选依赖，pyproject doc 组）
    except Exception as exc:
        raise ValueError(
            "docx 导出需要可选依赖 python-docx：py -X utf8 -m pip install python-docx") from exc

    p = Path(out_path)
    p.parent.mkdir(parents=True, exist_ok=True)
    anchors.attach_anchors(report)
    sanitize_paths(report)

    document = docx.Document()
    document.add_heading("安全风险隐患研判报告", level=0)

    document.add_paragraph("源文件：%s" % report.get("source_file", ""))
    document.add_paragraph("记录数：%d，识别隐患：%d" % (
        report.get("record_count", 0), report.get("hazard_count", 0)))
    document.add_paragraph("生成：SiteGuard v%s（LLM 兜底：%s）" % (
        report.get("version", "?"),
        "关闭" if report.get("settings", {}).get("disable_llm") else "开启"))

    cards = report.get("cards", [])
    if cards:
        document.add_heading("一、隐患清单", level=1)
        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        header = table.rows[0].cells
        for i, h in enumerate(("#", "原文摘录", "隐患类型", "风险等级", "依据条款", "整改建议")):
            header[i].text = h
        for i, card in enumerate(cards, 1):
            j = card.get("judgment", {})
            clauses = "；".join(
                "%s %s" % (r.get("standard", ""), r.get("clause", ""))
                for r in j.get("clause_refs", [])) or "无条款"
            sugg = "；".join(s["text"] for s in j.get("suggestions", [])) or "无"
            row = table.add_row().cells
            for k, v in enumerate((str(i), card.get("quote", ""),
                                   card.get("hazard_type", ""), j.get("risk_level", "—"),
                                   clauses, sugg)):
                row[k].text = v

    confirm = [c for c in cards if c.get("judgment", {}).get("human_confirm")]
    if confirm:
        document.add_heading("二、待人工确认", level=1)
        for c in confirm:
            document.add_paragraph("%s（%s）" % (c.get("quote", ""),
                                                c["judgment"]["hazard_type"]))

    if report.get("warnings"):
        document.add_heading("三、警告", level=1)
        for w in report["warnings"]:
            document.add_paragraph(w)

    document.add_paragraph(
        "本报告由规则引擎生成，风险等级与条款依据可溯源；种子条款以官方文本为准。"
        "条款依据可联动知识图谱可视化 kg.html（同目录部署时点击条款锚点定位条文节点）。")

    # 签署栏（M5）：编制/审核/批准 + 日期占位线
    document.add_heading("签署栏", level=1)
    sign = document.add_table(rows=3, cols=4)
    sign.style = "Table Grid"
    for r, role in enumerate(("编制人", "审核人", "批准人")):
        cells = sign.rows[r].cells
        cells[0].text = role
        cells[1].text = ""
        cells[2].text = "日期"
        cells[3].text = "____年__月__日"

    document.save(str(p))
    return p
