"""文本清洗与记录解析（规则优先通路的第一层）。

M1 v2：
- 时间戳容错：2026-09-21 / 2026/09/21 / 2026.09.21 / 2026年9月21日，
  时间 HH:MM / HH:MM:SS / HH时MM分（均可省略），统一规范化为 "YYYY-MM-DD HH:MM"；
- 多行记录合并：时间戳行开启新记录，其后不匹配时间戳、非注释/表头的行视为
  上一条记录的续行，按邻近合并（中日韩相邻直接拼接，ASCII 词间补空格）。

M2：docx 解析接入（python-docx 走 optional 依赖 pyproject doc 组）；
未安装依赖/坏文件抛 ValueError 带安装提示，绝不崩溃主流程。
"""

import re
from pathlib import Path

# 记录行格式：[日期 时间] 描述文本（日期/时间分隔符与格式容错）
TS_RE = (r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?\s*"
         r"(?:(\d{1,2}):(\d{1,2})(?::(\d{1,2}))?|(\d{1,2})时(\d{1,2})分?)?")
RECORD_RE = re.compile(r"^\[" + TS_RE + r"\]\s*(.+)$")

# 续行合并时跳过的注释/备注行
SKIP_PREFIXES = ("#", "//", "备注", "注：", "注:")

SUPPORTED_SUFFIXES = {".txt", ".docx"}


def _docx_text(path):
    """docx → 纯文本（段落 + 表格行）。缺依赖/坏文件抛 ValueError，绝不崩溃主流程。"""
    try:
        import docx  # python-docx（可选依赖，pyproject doc 组）
    except ImportError as exc:
        raise ValueError(
            "docx 解析需要可选依赖 python-docx：py -X utf8 -m pip install python-docx") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ValueError("docx 文件损坏或不是有效 Word 文档: %s（%s: %s）"
                         % (path, type(exc).__name__, exc)) from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def clean_text(raw):
    """清洗：零宽字符/BOM、CRLF、中文间伪空格、全角空格；去空行。

    中文间空格（PDF/导出文本常见逐字符伪空格）必须去除，否则关键词正则失配。
    """
    for ch in ("\u200b", "\u200c", "\u200d", "\ufeff"):
        raw = raw.replace(ch, "")
    raw = raw.replace("\r\n", "\n").replace("\r", "\n")
    raw = raw.replace("\u3000", " ")
    raw = re.sub(r"([\u4e00-\u9fff])[ \t]+(?=[\u4e00-\u9fff])", r"\1", raw)
    lines = [ln.strip() for ln in raw.split("\n")]
    return "\n".join(ln for ln in lines if ln)


def normalize_timestamp(match):
    """RECORD_RE 的匹配结果 → 规范化 "YYYY-MM-DD HH:MM"（无时间则仅日期）。"""
    year, month, day = int(match.group(1)), int(match.group(2)), int(match.group(3))
    hour = match.group(4) if match.group(4) is not None else match.group(7)
    minute = match.group(5) if match.group(5) is not None else match.group(8)
    ts = "%04d-%02d-%02d" % (year, month, day)
    if hour is not None and minute is not None:
        ts += " %02d:%02d" % (int(hour), int(minute))
    return ts


def _join_merged(left, right):
    """续行拼接：中日韩相邻直接拼接，两侧均为 ASCII 字母数字时补空格。"""
    if left and right and left[-1].isascii() and left[-1].isalnum() \
            and right[0].isascii() and right[0].isalnum():
        return left + " " + right
    return left + right


def parse_lines(text):
    """把清洗后的文本解析为记录列表。

    返回 [{index, timestamp, text}]；index 为记录序号（1 起），不是物理行号，
    保证生成器真值（gold）与解析结果对齐稳定。同一记录跨行时按邻近合并。
    """
    records = []
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            continue
        m = RECORD_RE.match(line)
        if m:
            records.append({"index": len(records) + 1,
                            "timestamp": normalize_timestamp(m),
                            "text": m.group(9).strip()})
            continue
        if line.startswith(SKIP_PREFIXES):
            continue
        if records:
            records[-1]["text"] = _join_merged(records[-1]["text"], line)
        # 首条记录之前的行视为表头，跳过
    return records


def parse_log_file(path):
    """解析单个日志文件（.txt/.docx）。坏格式/不支持的类型抛 ValueError，缺文件抛 FileNotFoundError。"""
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError("输入文件不存在: %s" % p)
    suffix = p.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError("暂不支持的文件类型: %s（支持 .txt/.docx）" % p.suffix)
    if suffix == ".docx":
        text = _docx_text(p)
    else:
        raw = p.read_bytes()
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            raise ValueError("文件不是有效 UTF-8（疑似 GBK 或二进制）: %s" % p)
    return parse_lines(clean_text(text))
