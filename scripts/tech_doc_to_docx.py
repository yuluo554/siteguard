# -*- coding: utf-8 -*-
"""技术方案 docx 固化（M6）：docs/技术方案.md → docs/技术方案.docx 程序化生成。

纪律（"提交材料固化"）：
- **禁止手工改后入库**：docx 只能由本脚本从 docs/技术方案.md 程序化生成；
  修改需求一律改 md 源后重跑本脚本，并同步更新 data/README.md 台账与
  scripts/desensitize_audit.py 二进制白名单中的 sha256（docx 为 zip 容器，
  重新生成的字节会因 zip 时间戳差异而变化）；
- python-docx 为可选依赖（pyproject ``doc`` 组），缺依赖抛 ValueError 带安装
  提示（照 report/export.py 范式），不崩溃；
- 版式：md 标题→docx 标题、表格→Table Grid 表格、围栏代码块（含 mermaid 图源）
  →等宽文本段落保留源码、列表/引用→对应样式；行内 **粗体** 与 `代码` 保留；
  不嵌入任何图片与超链接，docx 天然 0 外链（脱敏④步白名单管理）。

用法：py -X utf8 -m scripts.tech_doc_to_docx [--src docs/技术方案.md --out docs/技术方案.docx]
成功末行打印 sha256 供台账登记。
"""

import argparse
import hashlib
import re
import sys
from pathlib import Path

from siteguard import config

_INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
_LIST_RE = re.compile(r"^(\d+)\.\s+")


def _set_eastasia(run):
    """中文 run 补设 eastAsia 字体，避免等宽/西文字体下中文回退不一致。"""
    try:
        from docx.oxml.ns import qn
        rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
        rfonts.set(qn("w:eastAsia"), "宋体")
    except Exception:
        pass


def _add_inline(par, text, bold_all=False):
    """行内 **粗体** / `代码` / 纯文本 → runs。"""
    for part in _INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            run = par.add_run(part)
            if bold_all:
                run.bold = True


def _add_code_block(doc, lines):
    par = doc.add_paragraph()
    run = par.add_run("\n".join(lines))
    run.font.name = "Consolas"
    _set_eastasia(run)
    return par


def _add_table(doc, rows):
    grid = []
    for r in rows:
        cells = [c.strip() for c in r.strip().strip("|").split("|")]
        if cells and all(set(c) <= set("-: ") for c in cells):
            continue  # 表头分隔行
        grid.append(cells)
    if not grid:
        return
    n_cols = max(len(r) for r in grid)
    table = doc.add_table(rows=len(grid), cols=n_cols)
    table.style = "Table Grid"
    for ri, row in enumerate(grid):
        for ci in range(n_cols):
            cell = table.rows[ri].cells[ci]
            cell.text = ""
            _add_inline(cell.paragraphs[0], row[ci] if ci < len(row) else "",
                        bold_all=(ri == 0))


def _is_marker(line):
    """块起始标记：标题/围栏/表格/引用/列表项（列表续行不再并入新块）。"""
    s = line.strip()
    return (s.startswith("#") or s.startswith("```") or s.startswith("|")
            or s.startswith("> ") or s.startswith("- ")
            or bool(_LIST_RE.match(s)))


def _blocks(lines):
    """md 行流 → 块流。硬换行段落/列表项跨行合并为单块（CJK 直接拼接不加空格），
    使跨行的 **粗体**、`代码` 等行内标记能被完整解析。"""
    i, n = 0, len(lines)
    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1
            continue
        if s.startswith("```"):
            i += 1
            code = []
            while i < n and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1  # 跳过收尾 ```
            yield ("code", code)
            continue
        if s.startswith("|"):
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append(lines[i].strip())
                i += 1
            yield ("table", rows)
            continue
        if s.startswith("#"):
            yield ("heading", s)
            i += 1
            continue
        if s.startswith("> "):
            yield ("quote", s[2:])
            i += 1
            continue
        kind = "bullet" if s.startswith("- ") else (
            "numbered" if _LIST_RE.match(s) else "para")
        text = s[2:] if kind == "bullet" else s
        i += 1
        while i < n and lines[i].strip() and not _is_marker(lines[i]):
            text += lines[i].strip()
            i += 1
        yield (kind, text)


def convert(md_text, doc):
    """markdown 全文 → python-docx Document（就地追加）。"""
    for kind, payload in _blocks(md_text.splitlines()):
        if kind == "code":
            _add_code_block(doc, payload)
        elif kind == "table":
            _add_table(doc, payload)
        elif kind == "heading":
            if payload.startswith("### "):
                doc.add_heading(payload[4:].strip(), level=3)
            elif payload.startswith("## "):
                doc.add_heading(payload[3:].strip(), level=1)
            elif payload.startswith("# "):
                doc.add_heading(payload[2:].strip(), level=0)
        elif kind == "quote":
            par = doc.add_paragraph()
            _add_inline(par, payload)
            for run in par.runs:
                run.italic = True
        elif kind == "bullet":
            par = doc.add_paragraph(style="List Bullet")
            _add_inline(par, payload)
        else:  # para / numbered：编号保留字面文本，避免 docx 编号重启问题
            par = doc.add_paragraph()
            _add_inline(par, payload)


def main(argv=None):
    ap = argparse.ArgumentParser(
        prog="tech_doc_to_docx", description=__doc__.splitlines()[0])
    ap.add_argument("--src", default="docs/技术方案.md", help="markdown 源文件")
    ap.add_argument("--out", default="docs/技术方案.docx", help="docx 输出路径")
    args = ap.parse_args(argv)
    try:
        import docx  # noqa: F401  python-docx（可选依赖，pyproject doc 组）
    except Exception as exc:
        raise ValueError("docx 生成需要可选依赖 python-docx："
                         "py -X utf8 -m pip install python-docx") from exc

    src = config.repo_path(args.src)
    out = config.repo_path(args.out)
    if not src.exists():
        raise FileNotFoundError("源文件不存在：%s" % src)
    document = docx.Document()
    convert(src.read_text(encoding="utf-8"), document)
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    data = out.read_bytes()
    digest = hashlib.sha256(data).hexdigest()
    print("已生成 %s（%d 字节，程序化生成自 %s，禁止手工修改）" % (args.out, len(data), args.src))
    print("sha256：%s" % digest)
    return 0


if __name__ == "__main__":
    sys.exit(main())
