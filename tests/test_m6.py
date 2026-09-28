# -*- coding: utf-8 -*-
"""M6 收尾加分项测试：向量检索兜底、视频抽帧关键帧（可注入、评测零依赖）。

注入说明（照 M3 pyvis_mod / M4 detector / M5 serve 可注入范式）：
- fake embedder / fake reader / fake detector 仅为测试注入，用于验证通路与降级，
  不伪造任何真实检测结果；真实推理路径（opencv/ultralytics）在缺依赖环境走降级；
- 向量检索兜底为纯 stdlib 实现，测试直接用真实条文库构建索引（无 fake 成分）。
"""
import io
import json
import shutil
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

from siteguard import config
from siteguard.extract import video as video_mod
from siteguard.retrieval import vector as vector_mod

REPO = config.repo_path()


def _load_reg_count():
    path = config.repo_path("data", "knowledge", "regulations",
                            "seed_regulations.jsonl")
    return sum(1 for line in path.read_text(encoding="utf-8").splitlines()
               if line.strip())


class TestVectorIndex(unittest.TestCase):
    """词法向量索引：真实条文库建索引、排序有界、空索引、embedder 注入。"""

    def test_lexical_top1_hit(self):
        idx = vector_mod.build_regulation_index()
        self.assertEqual(len(idx.docs), _load_reg_count())
        hits = idx.search("临边作业 防护栏杆 密目式安全立网 封闭", top_k=3)
        self.assertEqual(len(hits), 3)
        top, score = hits[0]
        self.assertEqual(top["standard"], "JGJ 80-2016")
        self.assertEqual(top["clause"], "4.1.2")
        self.assertGreater(score, 0.0)

    def test_scores_sorted_and_bounded(self):
        idx = vector_mod.build_regulation_index()
        hits = idx.search("配电箱 箱门 配锁 专人", top_k=5)
        scores = [s for _, s in hits]
        self.assertEqual(scores, sorted(scores, reverse=True))
        self.assertTrue(all(s >= 0.0 for s in scores))

    def test_empty_index_returns_empty(self):
        idx = vector_mod.VectorIndex([])
        self.assertEqual(idx.search("任何查询"), [])

    def test_injected_embedder_dense_path(self):
        # fake embedder 仅测试注入：固定 3 维向量表，验证稠密向量通路
        table = {"安全带": [1.0, 0.0, 0.0], "配电箱": [0.0, 1.0, 0.0],
                 "消防": [0.0, 0.0, 1.0]}

        def fake_embedder(text):
            for k, v in table.items():
                if k in str(text):
                    return v
            return [0.0, 0.0, 0.0]

        docs = [
            {"id": "a", "text": "安全带 高挂低用", "standard": "S", "clause": "1",
             "std_name": ""},
            {"id": "b", "text": "配电箱 箱门配锁", "standard": "S", "clause": "2",
             "std_name": ""},
        ]
        idx = vector_mod.VectorIndex(docs, embedder=fake_embedder)
        hits = idx.search("安全带", top_k=1)
        self.assertEqual(hits[0][0]["id"], "a")
        self.assertGreater(hits[0][1], 0.99)


class TestVectorFallback(unittest.TestCase):
    """兜底候选：仅补空条款卡片、判定字段零改动、MD 导出带提示段。"""

    @staticmethod
    def _report(cards):
        return {"source_file": "x.txt", "record_count": 1,
                "hazard_count": len(cards), "cards": cards, "warnings": []}

    def test_fallback_attaches_only_to_empty_clause_cards(self):
        idx = vector_mod.build_regulation_index()
        empty = {"quote": "外脚手架临边未设置防护栏杆也未挂安全网",
                 "hazard_type": "临边防护缺失",
                 "judgment": {"risk_level": "较大风险", "clause_refs": []}}
        linked = {"quote": "配电箱箱门未配锁", "hazard_type": "配电箱缺陷",
                  "judgment": {"risk_level": "一般风险",
                               "clause_refs": [{"standard": "JGJ 46-2005",
                                                "clause": "8.3.2"}]}}
        report = self._report([empty, linked])
        n = vector_mod.fallback_clauses(report, idx, top_k=2)
        self.assertEqual(n, 1)
        fb = empty["fallback_clauses"]
        self.assertTrue(fb)
        for x in fb:
            self.assertTrue(x["standard"] and x["clause"])
            self.assertIn("非判定依据", x["note"])
            self.assertGreaterEqual(x["score"], 0.05)
        self.assertNotIn("fallback_clauses", linked)
        self.assertEqual(linked["judgment"]["clause_refs"][0]["clause"], "8.3.2")
        self.assertEqual(empty["judgment"]["clause_refs"], [])

    def test_md_export_fallback_section(self):
        from siteguard.report import export
        card = {"quote": "临边未防护", "hazard_type": "临边防护缺失",
                "category": "高处作业",
                "fallback_clauses": [{"standard": "JGJ 80-2016", "clause": "4.1.2",
                                      "std_name": "", "score": 0.42,
                                      "note": vector_mod.FALLBACK_NOTE}],
                "judgment": {"risk_level": "较大风险", "clause_refs": [],
                             "suggestions": []}}
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        p = export.export_md(self._report([card]), tmp / "r.md")
        text = p.read_text(encoding="utf-8")
        self.assertIn("兜底候选条款", text)
        self.assertIn("非判定依据", text)
        self.assertIn("JGJ 80-2016 4.1.2", text)


class TestVideoKeyframes(unittest.TestCase):
    """视频抽帧：纯函数选帧、注入 reader 落盘、缺 opencv 降级、注入全通路。"""

    def test_select_keyframes_pure_function(self):
        self.assertEqual(video_mod.select_keyframes(0), [])
        self.assertEqual(video_mod.select_keyframes(-3), [])
        self.assertEqual(video_mod.select_keyframes(5, 8), [0, 1, 2, 3, 4])
        ks = video_mod.select_keyframes(100, 8)
        self.assertEqual(len(ks), 8)
        self.assertEqual(ks[0], 0)
        self.assertTrue(all(b > a for a, b in zip(ks, ks[1:])))
        self.assertLess(ks[-1], 100)

    def test_extract_with_injected_reader(self):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        total, max_frames = 25, 8

        def fake_reader(_path):
            # fake reader 仅测试注入：合成以 JPEG 魔数开头的字节，非真实视频帧
            def get_frame(idx):
                return b"\xff\xd8\xffFAKE" + bytes([idx % 256]) * 8
            return total, get_frame

        paths = video_mod.extract_keyframes("fake.mp4", tmp / "frames",
                                            max_frames=max_frames,
                                            reader=fake_reader)
        expect = video_mod.select_keyframes(total, max_frames)
        self.assertEqual(len(paths), len(expect))
        for p in paths:
            self.assertTrue(p.exists())
            self.assertEqual(p.read_bytes()[:3], b"\xff\xd8\xff")

    def test_missing_opencv_valueerror_with_hint(self):
        with mock.patch.object(video_mod, "_import_cv2",
                               side_effect=ValueError(video_mod.INSTALL_HINT)):
            tmp = Path(tempfile.mkdtemp())
            self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
            with self.assertRaises(ValueError) as ctx:
                video_mod.extract_keyframes("x.mp4", tmp / "f")
        self.assertIn("opencv-python", str(ctx.exception))

    def test_collect_video_cards_with_injections(self):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)

        def fake_reader(_path):
            # fake reader 仅测试注入：10 帧合成图像字节
            return 10, (lambda idx: b"\xff\xd8\xffFAKE" + bytes([idx]) * 8)

        def fake_detector(_image_path):
            # fake detector 仅测试注入：稳定返回"未佩戴安全帽"检出（head），
            # 映射受控词表 HT-GCZY-007 走真实 schema 校验
            return [{"class_name": "head", "bbox": [1.0, 2.0, 3.0, 4.0],
                     "confidence": 0.9}]

        cards, warnings = video_mod.collect_video_cards(
            "fake.mp4", tmp / "frames", max_frames=4,
            detector=fake_detector, reader=fake_reader)
        self.assertEqual(len(cards), 4)
        for c in cards:
            self.assertEqual(c["extractor"], "vision")
            self.assertEqual(c["hazard_type_id"], "HT-GCZY-007")
        self.assertEqual(warnings, [])


class TestCliM6Flags(unittest.TestCase):
    """demo 新旗标：--vector-fallback 生效、--video 缺依赖/坏文件降级 exit 0。"""

    def test_demo_vector_fallback_flag(self):
        from siteguard.cli import main
        tmp = "examples/m6_cli_tmp"
        self.addCleanup(shutil.rmtree, config.repo_path(tmp), ignore_errors=True)
        buf = io.StringIO()
        with redirect_stdout(buf):
            rc = main(["demo", "--scenario", "gczy", "--vector-fallback",
                       "--out", tmp])
        self.assertEqual(rc, 0)
        self.assertIn("向量检索兜底", buf.getvalue())

    def test_demo_video_degrades_gracefully(self):
        # 缺 opencv 或视频文件无效两态均打印降级说明且 exit 0（本机缺 opencv 实测前者）
        from siteguard.cli import main
        tmp = "examples/m6_cli_tmp2"
        self.addCleanup(shutil.rmtree, config.repo_path(tmp), ignore_errors=True)
        buf = io.StringIO()
        with redirect_stdout(buf):
            rc = main(["demo", "--scenario", "gczy",
                       "--video", "no_such_file.mp4", "--out", tmp])
        self.assertEqual(rc, 0)
        self.assertIn("视频抽帧降级", buf.getvalue())


class TestTechDocDocx(unittest.TestCase):
    """技术方案 docx 固化物守门：程序化生成、结构完整、0 外链。"""

    def test_docx_structure_and_zero_external(self):
        path = REPO / "docs" / "技术方案.docx"
        self.assertTrue(path.exists(), "docs/技术方案.docx 必须程序化生成后入仓")
        try:
            import docx
        except Exception:
            self.skipTest("缺 python-docx（optional doc 组）")
        d = docx.Document(str(path))
        h1 = [p.text for p in d.paragraphs
              if p.style.name.startswith("Heading 1")]
        self.assertEqual(len(h1), 9)
        self.assertIn("9. 创新点小结", h1)
        self.assertEqual(len(d.tables), 2)
        all_text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("向量检索兜底", all_text)
        self.assertIn("视频抽帧", all_text)
        with __import__("zipfile").ZipFile(str(path)) as z:
            external = [n for n in z.namelist()
                        if n.endswith(".rels") and b"External" in z.read(n)]
        self.assertEqual(external, [])


if __name__ == "__main__":
    unittest.main()
