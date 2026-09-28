# -*- coding: utf-8 -*-
"""M5 交付打磨测试：export_docx/签署栏、路径脱敏、Web 面板 serve、脱敏审查守门。

缺依赖路径通过注入模拟（照 M3 pyvis_mod / M4 detector 可注入范式）：
- export_docx 缺 python-docx：patch builtins.__import__ 使 import docx 失败；
- serve 缺 fastapi：monkeypatch serverapp._import_fastapi 抛 ImportError；
fake 仅为模拟"缺依赖"，不伪造任何研判/检测结果。
"""
import io
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from pathlib import Path

from siteguard import config
from siteguard.report import export
from siteguard.serverapp import INSTALL_HINT, analyze_input, create_app, scenario_payload

REPO = config.repo_path()


def _run_demo_report():
    """跑一次高处作业样例研判（真实规则通路，非 fake），供导出测试用。"""
    from siteguard.pipeline import runner
    return runner.run_file(config.repo_path("data", "samples", "text", "demo_gczy.txt"))


class TestExportDocx(unittest.TestCase):
    """docx 导出：签署栏齐全、0 外链图片、缺依赖 ValueError 带安装提示。"""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)
        self.report = _run_demo_report()

    def test_docx_with_signature_block(self):
        p = export.export_docx(self.report, self.tmp / "deep" / "sub" / "t.docx")
        self.assertTrue(p.exists())
        import docx
        d = docx.Document(str(p))
        texts = [x.text for x in d.paragraphs]
        texts += [c.text for t in d.tables for row in t.rows for c in row.cells]
        self.assertTrue(any("签署栏" in x for x in texts))
        for role in ("编制人", "审核人", "批准人"):
            self.assertTrue(any(role in x for x in texts), role)
        self.assertTrue(any("____年__月__日" in x for x in texts))
        self.assertTrue(any("隐患清单" in x for x in texts))

    def test_docx_no_external_image_links(self):
        p = export.export_docx(self.report, self.tmp / "t.docx")
        with zipfile.ZipFile(str(p)) as z:
            external = [n for n in z.namelist()
                        if n.endswith(".rels") and b"External" in z.read(n)]
        self.assertEqual(external, [])

    def test_docx_missing_dependency_valueerror(self):
        import builtins
        orig = builtins.__import__

        def fake_import(name, *a, **kw):
            if name == "docx":
                raise ImportError("No module named 'docx'")
            return orig(name, *a, **kw)

        builtins.__import__ = fake_import
        try:
            with self.assertRaises(ValueError) as ctx:
                export.export_docx(self.report, self.tmp / "t.docx")
        finally:
            builtins.__import__ = orig
        self.assertIn("python-docx", str(ctx.exception))
        self.assertIn("pip install", str(ctx.exception))


class TestSanitizePaths(unittest.TestCase):
    """脱敏第③步：仓库内路径 → 相对 posix；导出产物不含本机绝对路径。"""

    def test_display_path_inside_repo(self):
        p = config.repo_path("data", "samples", "text", "demo_gczy.txt")
        self.assertEqual(config.display_path(p), "data/samples/text/demo_gczy.txt")

    def test_display_path_outside_repo_untouched(self):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        self.assertEqual(config.display_path(tmp / "x.txt"), str(tmp / "x.txt"))

    def test_export_json_artifact_relative(self):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, ignore_errors=True)
        report = _run_demo_report()
        export.export_json(report, tmp / "r.json")
        payload = json.loads((tmp / "r.json").read_text(encoding="utf-8"))
        self.assertEqual(payload["source_file"], "data/samples/text/demo_gczy.txt")
        for card in payload["cards"]:
            self.assertFalse(card["source"]["file"].startswith("D:"))
        self.assertNotIn("ProgramData", (tmp / "r.json").read_text(encoding="utf-8"))

    def test_committed_demo_artifacts_relative(self):
        """入库产物（demo/eval 报告）不含本机路径——脱敏审查守门。"""
        for rel in ("examples/demo_output/demo_gczy.report.json",
                    "examples/demo_output/demo_lsyd.report.json",
                    "examples/demo_output/demo_gczy.report.md",
                    "examples/eval_report.json"):
            raw = (REPO / rel).read_text(encoding="utf-8")
            self.assertNotIn("ProgramData", raw, rel)
            self.assertNotIn("AppData", raw, rel)


class TestServeCore(unittest.TestCase):
    """serve 核心纯函数与缺依赖降级（不依赖 fastapi 可用性）。"""

    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, ignore_errors=True)

    def test_scenario_payload_all_and_single(self):
        self.assertEqual([x["key"] for x in scenario_payload("all")], ["gczy", "lsyd"])
        self.assertEqual([x["key"] for x in scenario_payload("gczy")], ["gczy"])
        with self.assertRaises(ValueError):
            scenario_payload("nope")

    def test_analyze_input_two_scenarios(self):
        for key in ("gczy", "lsyd"):
            txt = config.repo_path("data", "samples", "text", "demo_%s.txt" % key)
            out = analyze_input(txt.read_bytes(), txt.name, self.tmp)
            self.assertEqual(out["report"]["hazard_count"], 54, key)
            self.assertEqual(out["report"]["source_file"], txt.name)
            self.assertTrue(out["downloads"]["docx"])
            self.assertTrue((self.tmp / ("demo_%s.report.json" % key)).exists())

    def test_analyze_input_rejects_bad_suffix_and_empty(self):
        with self.assertRaises(ValueError):
            analyze_input(b"x=1", "evil.exe", self.tmp)
        with self.assertRaises(ValueError):
            analyze_input("   ", "pasted.txt", self.tmp)

    def test_analyze_input_filename_traversal_cleaned(self):
        txt = config.repo_path("data", "samples", "text", "demo_gczy.txt")
        out = analyze_input(txt.read_bytes(), "..\\..\\evil\\demo_gczy.txt", self.tmp)
        self.assertEqual(out["report"]["source_file"], "demo_gczy.txt")

    def test_analyze_input_bad_docx_degrades(self):
        out = analyze_input(b"not a real docx", "bad.docx", self.tmp)
        self.assertEqual(out["report"]["hazard_count"], 0)
        self.assertTrue(any("docx" in w for w in out["report"]["warnings"]))

    def test_create_app_missing_fastapi_runtimeerror(self):
        from siteguard import serverapp
        old = serverapp._import_fastapi
        serverapp._import_fastapi = lambda: (_ for _ in ()).throw(
            ImportError("No module named 'fastapi'"))
        try:
            with self.assertRaises(RuntimeError) as ctx:
                create_app()
        finally:
            serverapp._import_fastapi = old
        self.assertIn("pip install fastapi uvicorn", str(ctx.exception))

    def test_cmd_serve_missing_deps_exit2(self):
        from siteguard import serverapp
        from siteguard.cli import main
        old = serverapp._import_fastapi
        serverapp._import_fastapi = lambda: (_ for _ in ()).throw(
            ImportError("No module named 'fastapi'"))
        try:
            buf = io.StringIO()
            with redirect_stdout(buf):
                rc = main(["serve"])
        finally:
            serverapp._import_fastapi = old
        self.assertEqual(rc, 2)
        self.assertIn("pip install fastapi uvicorn", buf.getvalue())


class TestServeHTTP(unittest.TestCase):
    """HTTP 层（fastapi TestClient）：两场景研判 + 失败路径 + 场景覆盖 + 0 外链。"""

    @classmethod
    def setUpClass(cls):
        from fastapi.testclient import TestClient
        cls.client = TestClient(create_app())

    def test_index_and_static_no_external_links(self):
        for path in ("/", "/static/app.js", "/static/style.css"):
            r = self.client.get(path)
            self.assertEqual(r.status_code == 200, True, path)
            self.assertNotIn("http://", r.text, path)
            self.assertNotIn("https://", r.text, path)

    def test_analyze_two_scenarios_real_pipeline(self):
        for key in ("gczy", "lsyd"):
            s = self.client.get("/api/sample/" + key).json()
            r = self.client.post("/api/analyze",
                                 data={"text": s["text"], "filename": s["filename"]})
            self.assertEqual(r.status_code, 200)
            body = r.json()
            self.assertEqual(body["hazard_count"], 54, key)
            self.assertTrue(body["downloads"]["docx"])
            md = self.client.get("/api/report/%s/md" % body["report_id"])
            self.assertEqual(md.status_code, 200)
            self.assertIn("研判报告", md.text)

    def test_failure_paths(self):
        self.assertEqual(
            self.client.post("/api/analyze",
                             data={"text": "x", "filename": "evil.exe"}).status_code, 400)
        self.assertEqual(self.client.post("/api/analyze", data={"text": " "}).status_code, 400)
        r = self.client.post("/api/analyze", files={"file": ("bad.docx", b"junk")})
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["hazard_count"], 0)
        self.assertTrue(r.json()["warnings"])
        self.assertEqual(self.client.get("/api/report/x/xml").status_code, 404)
        self.assertEqual(self.client.get("/api/report/no_such/md").status_code, 404)
        self.assertEqual(self.client.get("/api/sample/nope").status_code, 404)

    def test_scenario_override(self):
        from fastapi.testclient import TestClient
        c2 = TestClient(create_app(scenario="gczy"))
        self.assertEqual([x["key"] for x in c2.get("/api/scenarios").json()], ["gczy"])
        self.assertEqual(c2.get("/api/sample/lsyd").status_code, 404)
        self.assertEqual(c2.get("/api/sample/gczy").status_code, 200)

    def test_webapp_files_zero_external_links(self):
        for name in ("index.html", "app.js", "style.css"):
            raw = (REPO / "webapp" / name).read_text(encoding="utf-8")
            self.assertNotIn("http://", raw, name)
            self.assertNotIn("https://", raw, name)


class TestDesensitizeAuditGate(unittest.TestCase):
    """脱敏四步审查脚本守门：随测试重跑，四步全 PASS 才放行。"""

    def test_audit_pass(self):
        r = subprocess.run(
            [sys.executable, "-X", "utf8", "-B", "-m", "scripts.desensitize_audit"],
            capture_output=True, cwd=str(REPO))
        out = r.stdout.decode("utf-8", errors="replace")
        self.assertEqual(r.returncode, 0, out[-2000:])
        self.assertIn("DESENSITIZE_AUDIT_OK", out)


if __name__ == "__main__":
    unittest.main()
